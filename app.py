
import os
import streamlit as st

from dotenv import load_dotenv
from PyPDF2 import PdfReader
from docx import Document

from langchain_openai import ChatOpenAI
from langchain_core.prompts import ChatPromptTemplate



# OpenAI API Key from Streamlit Secrets
try:
    openai_api_key = st.secrets["OPENAI_API_KEY"]
except KeyError:
    st.error("OPENAI_API_KEY is not configured in Streamlit Secrets.")
    st.stop()








st.set_page_config(
    page_title="AI Resume Agent",
    page_icon="🤖",
    layout="wide"
)


# ============================================================
# OPENAI MODEL
# ============================================================

llm = ChatOpenAI(
    model="gpt-4o-mini",
    temperature=0.4,
    api_key=openai_api_key
)


# ============================================================
# SESSION STATE
# ============================================================

if "resume_text" not in st.session_state:
    st.session_state.resume_text = None

if "resume_name" not in st.session_state:
    st.session_state.resume_name = None

if "chat_history" not in st.session_state:
    st.session_state.chat_history = []


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file):

    pdf_reader = PdfReader(file)

    text = ""

    for page in pdf_reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_docx_text(file):

    document = Document(file)

    text = ""

    # Paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text + "\n"

    # Tables
    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():
                    text += cell.text + "\n"

    return text


# ============================================================
# RESUME EXTRACTION
# ============================================================

def extract_resume(file):

    file_name = file.name.lower()

    if file_name.endswith(".pdf"):

        return extract_pdf_text(file)

    elif file_name.endswith(".docx"):

        return extract_docx_text(file)

    else:

        return None


# ============================================================
# GENERAL AI
# ============================================================

def general_chat(question):

    prompt = ChatPromptTemplate.from_messages([

        (
            "system",
            """
            You are a helpful AI assistant.

            Answer the user's question clearly and accurately.

            If the question is technical:
            - Explain it clearly
            - Give examples when useful
            - Use simple language when appropriate

            Do not unnecessarily mention resumes.
            """
        ),

        (
            "human",
            "{question}"
        )

    ])

    chain = prompt | llm

    response = chain.invoke({
        "question": question
    })

    return response.content


# ============================================================
# RESUME QUESTION
# ============================================================

def ask_resume_question(resume, question):

    prompt = ChatPromptTemplate.from_messages([

        (
            "system",
            """
            You are an expert Resume AI Agent.

            The user has uploaded a resume.

            Answer questions about the user's resume
            ONLY using information available in the resume.

            If the requested information is not available,
            clearly say:

            "This information is not mentioned in the resume."

            Do not invent:
            - Experience
            - Skills
            - Companies
            - Certifications
            - Education
            - Projects
            - Achievements

            Resume:

            {resume}
            """
        ),

        (
            "human",
            "{question}"
        )

    ])

    chain = prompt | llm

    response = chain.invoke({

        "resume": resume,
        "question": question

    })

    return response.content


# ============================================================
# RESUME ANALYSIS
# ============================================================

def analyze_resume(resume):

    prompt = ChatPromptTemplate.from_messages([

        (
            "system",
            """
            You are a professional resume reviewer and ATS expert.

            Analyze the following resume.

            Provide the following sections:

            1. Overall Resume Quality
            2. Estimated ATS Score / 100
            3. Strong Points
            4. Weak Points
            5. Missing Skills
            6. Missing Keywords
            7. Formatting Issues
            8. Experience Section Improvements
            9. Professional Summary Improvements
            10. Education Improvements
            11. Project Improvements
            12. Actionable Recommendations

            IMPORTANT:

            Do not invent experience, skills,
            education, projects or achievements.

            The ATS score is an estimated score,
            not a score from an actual ATS system.

            Resume:

            {resume}
            """
        )

    ])

    chain = prompt | llm

    response = chain.invoke({

        "resume": resume

    })

    return response.content


# ============================================================
# RESUME REFINEMENT
# ============================================================

def refine_resume(resume):

    prompt = ChatPromptTemplate.from_messages([

        (
            "system",
            """
            You are an expert Resume Writer and
            ATS Optimization Agent.

            Rewrite the user's resume to make it:

            - ATS friendly
            - Professional
            - Concise
            - Achievement focused
            - Keyword optimized
            - Easy to read
            - Recruiter friendly

            IMPORTANT:

            Do NOT invent:

            - Jobs
            - Companies
            - Skills
            - Certifications
            - Education
            - Projects
            - Achievements
            - Years of experience

            Only improve wording, clarity and structure
            using information already present.

            Keep all factual information truthful.

            Return the complete improved resume
            in a clean professional format.

            Resume:

            {resume}
            """
        )

    ])

    chain = prompt | llm

    response = chain.invoke({

        "resume": resume

    })

    return response.content


# ============================================================
# ATS OPTIMIZATION
# ============================================================

def optimize_resume_for_ats(resume, job_description):

    prompt = ChatPromptTemplate.from_messages([

        (
            "system",
            """
            You are an expert ATS Resume Optimization Agent.

            Compare the candidate's resume with the
            provided job description.

            Provide:

            1. ATS Match Score / 100
            2. Matching Keywords
            3. Missing Keywords
            4. Matching Skills
            5. Missing Skills
            6. Recommended Resume Changes
            7. Improved Professional Summary
            8. Improved Experience Bullets
            9. Optimized Skills Section
            10. Final Optimized Resume

            IMPORTANT:

            Never invent:

            - Experience
            - Companies
            - Skills
            - Certifications
            - Education
            - Projects
            - Achievements
            - Years of experience

            You may improve wording and highlight
            existing skills that match the job description.

            Do not add a skill unless it already exists
            in the candidate's resume.

            Resume:

            {resume}

            Job Description:

            {job_description}
            """
        )

    ])

    chain = prompt | llm

    response = chain.invoke({

        "resume": resume,
        "job_description": job_description

    })

    return response.content


# ============================================================
# JOB ROLE SUGGESTIONS
# ============================================================

def suggest_jobs(resume):

    prompt = ChatPromptTemplate.from_messages([

        (
            "system",
            """
            You are an expert career advisor.

            Analyze the resume and suggest suitable
            job roles.

            For each recommended role provide:

            1. Job Title
            2. Match Percentage
            3. Why the candidate matches
            4. Existing relevant skills
            5. Missing skills
            6. Recommended learning areas

            Rank the roles from strongest match
            to weakest match.

            Do not invent experience.

            Resume:

            {resume}
            """
        )

    ])

    chain = prompt | llm

    response = chain.invoke({

        "resume": resume

    })

    return response.content


# ============================================================
# INTERVIEW QUESTIONS
# ============================================================

def generate_interview_questions(resume):

    prompt = ChatPromptTemplate.from_messages([

        (
            "system",
            """
            You are an expert technical interviewer.

            Based ONLY on the candidate's resume,
            generate interview questions.

            Include:

            1. HR Questions
            2. Technical Questions
            3. Project Questions
            4. Previous Experience Questions
            5. Difficult Follow-up Questions
            6. Scenario-based Questions

            Do not ask about technologies,
            projects or experience that are not
            mentioned in the resume.

            Resume:

            {resume}
            """
        )

    ])

    chain = prompt | llm

    response = chain.invoke({

        "resume": resume

    })

    return response.content


# ============================================================
# AI REQUEST ROUTER
# ============================================================

def classify_request(question):

    prompt = ChatPromptTemplate.from_messages([

        (
            "system",
            """
            You are an AI request classifier.

            The user has uploaded a resume.

            Classify the user's request into EXACTLY
            ONE of these categories:

            GENERAL
            RESUME_QUESTION
            ANALYZE_RESUME
            REFINE_RESUME
            ATS_OPTIMIZATION
            JOB_ROLES
            INTERVIEW_QUESTIONS

            Rules:

            GENERAL
            Use when the user asks a general question
            unrelated to their resume.

            RESUME_QUESTION
            Use when the user asks for information
            about their own resume.

            Example:
            "What technologies are mentioned in my resume?"

            ANALYZE_RESUME
            Use when the user wants:
            - Resume analysis
            - Resume review
            - Strengths
            - Weaknesses
            - ATS score
            - Recommendations

            REFINE_RESUME
            Use when the user wants:
            - Resume rewriting
            - Resume improvement
            - Better wording
            - Professional rewrite

            ATS_OPTIMIZATION
            Use when the user wants:
            - ATS optimization
            - Job-specific resume optimization
            - Keyword matching
            - Resume matching with a job description

            JOB_ROLES
            Use when the user asks:
            - Which jobs suit me?
            - Suitable job roles
            - Career recommendations

            INTERVIEW_QUESTIONS
            Use when the user asks:
            - Interview questions
            - Technical interview preparation
            - HR questions
            - Mock interview questions

            Return ONLY the category name.

            User question:

            {question}
            """
        ),

        (
            "human",
            "{question}"
        )

    ])

    chain = prompt | llm

    response = chain.invoke({

        "question": question

    })

    category = response.content.strip().upper()

    return category




st.title(" AI Resume & General Chatbot")

st.write(
    "Ask anything, or upload your resume and ask the AI "
    "to analyze, refine, optimize, or answer questions "
    "about your resume."
)



with st.sidebar:

    st.header("📄 Resume")

    uploaded_file = st.file_uploader(
        "Upload Resume",
        type=["pdf", "docx"],
        help="Upload PDF or DOCX resume"
    )

    if uploaded_file:

        resume_text = extract_resume(uploaded_file)

        if resume_text:

            st.session_state.resume_text = resume_text

            st.session_state.resume_name = uploaded_file.name

            st.success(
                f"Resume loaded:\n{uploaded_file.name}"
            )

        else:

            st.error(
                "Unable to extract text from the resume."
            )

    elif st.session_state.resume_text:

        st.success(
            f"Resume loaded:\n"
            f"{st.session_state.resume_name}"
        )

    else:

        st.info(
            "No resume uploaded.\n\n"
            "You can still use the chatbot for "
            "general questions."
        )

    st.divider()

    if st.session_state.resume_text:

        if st.button(
            "🗑️ Remove Resume",
            use_container_width=True
        ):

            st.session_state.resume_text = None

            st.session_state.resume_name = None

            st.rerun()



if st.session_state.resume_text:

    with st.expander("📃 View Uploaded Resume"):

        st.text(
            st.session_state.resume_text
        )



for message in st.session_state.chat_history:

    with st.chat_message(message["role"]):

        if message.get("type") == "resume":

            st.text_area(
                message.get("title", "Result"),
                message["content"],
                height=500,
                key=message["key"]
            )

        else:

            st.markdown(
                message["content"]
            )



question = st.chat_input(
    "Ask anything or ask about your resume..."
)



if question:

    
    st.session_state.chat_history.append({

        "role": "user",
        "content": question

    })

    with st.chat_message("user"):

        st.markdown(question)


  
    with st.chat_message("assistant"):

        with st.spinner("🤖 Thinking..."):

            resume = st.session_state.resume_text


             

            if not resume:

                answer = general_chat(
                    question
                )

                st.markdown(answer)

                st.session_state.chat_history.append({

                    "role": "assistant",
                    "content": answer

                })


           
            else:

                category = classify_request(
                    question
                )


                
                if category == "GENERAL":

                    answer = general_chat(
                        question
                    )

                    st.markdown(answer)

                    st.session_state.chat_history.append({

                        "role": "assistant",
                        "content": answer

                    })


               
                elif category == "RESUME_QUESTION":

                    answer = ask_resume_question(
                        resume,
                        question
                    )

                    st.markdown(answer)

                    st.session_state.chat_history.append({

                        "role": "assistant",
                        "content": answer

                    })


              
                elif category == "ANALYZE_RESUME":

                    result = analyze_resume(
                        resume
                    )

                    st.markdown(result)

                    st.session_state.chat_history.append({

                        "role": "assistant",
                        "content": result

                    })


               
                elif category == "REFINE_RESUME":

                    result = refine_resume(
                        resume
                    )

                    st.text_area(
                        "✨ Refined Resume",
                        result,
                        height=600,
                        key=f"refined_{len(st.session_state.chat_history)}"
                    )

                    st.session_state.chat_history.append({

                        "role": "assistant",
                        "content": result,
                        "type": "resume",
                        "title": "✨ Refined Resume",
                        "key": f"history_refined_{len(st.session_state.chat_history)}"

                    })


               
                elif category == "ATS_OPTIMIZATION":

                    st.info(
                        "For the most accurate ATS optimization, "
                        "include the complete job description "
                        "in your message."
                    )

                    # The same chat message is used as the
                    # job description/request.

                    result = optimize_resume_for_ats(
                        resume,
                        question
                    )

                    st.text_area(
                        "🎯 ATS Optimized Resume",
                        result,
                        height=800,
                        key=f"ats_{len(st.session_state.chat_history)}"
                    )

                    st.session_state.chat_history.append({

                        "role": "assistant",
                        "content": result,
                        "type": "resume",
                        "title": "🎯 ATS Optimized Resume",
                        "key": f"history_ats_{len(st.session_state.chat_history)}"

                    })


               
                elif category == "JOB_ROLES":

                    result = suggest_jobs(
                        resume
                    )

                    st.markdown(result)

                    st.session_state.chat_history.append({

                        "role": "assistant",
                        "content": result

                    })


               
                elif category == "INTERVIEW_QUESTIONS":

                    result = generate_interview_questions(
                        resume
                    )

                    st.markdown(result)

                    st.session_state.chat_history.append({

                        "role": "assistant",
                        "content": result

                    })


               
                else:

                    answer = general_chat(
                        question
                    )

                    st.markdown(answer)

                    st.session_state.chat_history.append({

                        "role": "assistant",
                        "content": answer

                    })



if not st.session_state.chat_history:

    st.markdown("###  Try asking")

    col1, col2, col3 = st.columns(3)

   